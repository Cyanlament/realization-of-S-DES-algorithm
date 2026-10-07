"""Generate the Word and Markdown reports from experiment records."""
import argparse
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from report_content import report_pages

ROOT = Path(__file__).resolve().parents[1]


def set_font(style, family, size, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), family)


def make_document():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.4)
    section.left_margin = section.right_margin = Cm(2.5)
    section.header_distance, section.footer_distance = Cm(1.5), Cm(1.4)
    normal = document.styles["Normal"]
    set_font(normal, "宋体", 11)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 24), ("Subtitle", 16), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[name]
        set_font(style, "黑体", size, name.startswith("Heading"))
        style.paragraph_format.space_before = Pt(10 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
    for style in document.styles:
        for border in style.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
    set_font(document.styles["Caption"], "宋体", 10)
    document.styles["Caption"].paragraph_format.space_after = Pt(8)
    document.styles["Caption"].paragraph_format.line_spacing = 1.1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    properties = document.core_properties
    properties.title = "S-DES 算法实现实验报告"
    properties.subject = "信息安全导论 作业1"
    properties.author = "蔡旭涛；扶满"
    properties.last_modified_by = properties.comments = ""
    return document


def add_table(document, rows, widths):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "999999")):
            element.set(qn("w:" + key), value)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                element = OxmlElement("w:" + side)
                element.set(qn("w:w"), "85")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            properties.append(margins)
            if index == 0:
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "F0F0F0")
                properties.append(fill)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(str(value))
            run.font.size = Pt(10)
            run.bold = index == 0
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)
    spacer.add_run().font.size = Pt(4)


def write_reports(pages):
    document = make_document()
    markdown = []
    for page_index, elements in enumerate(pages):
        for element_index, element in enumerate(elements):
            kind, value = element[0], element[1]
            if kind in ("title", "subtitle", "h1", "h2", "h3"):
                style = {"title": "Title", "subtitle": "Subtitle", "h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[kind]
                paragraph = document.add_paragraph(value, style)
                if page_index and element_index == 0:
                    paragraph.paragraph_format.page_break_before = True
                if kind in ("title", "subtitle"):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                prefix = "#" * {"title": 1, "subtitle": 2, "h1": 2, "h2": 3, "h3": 4}[kind]
                markdown.append(prefix + " " + value)
            elif kind == "p":
                paragraph = document.add_paragraph(value)
                paragraph.paragraph_format.first_line_indent = Pt(22)
                markdown.append(value)
            elif kind == "table":
                add_table(document, value, element[2])
                lines = ["| " + " | ".join(str(cell).replace("\n", "<br>") for cell in row) + " |" for row in value]
                lines.insert(1, "| " + " | ".join("---" for _ in value[0]) + " |")
                markdown.append("\n".join(lines))
            elif kind == "image":
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = True
                shape = paragraph.add_run().add_picture(str(ROOT / value), width=Cm(element[3]))
                shape._inline.docPr.set("descr", element[2])
                caption = document.add_paragraph(element[2], "Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                markdown.append(f"![{element[2]}](../{value})")
            elif kind == "code":
                for line in value.splitlines():
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    run = paragraph.add_run(line)
                    run.font.name = "Consolas"
                    run.font.size = Pt(10)
                markdown.append("```text\n" + value + "\n```")
            elif kind == "gif":
                document.add_paragraph("演示动图：" + value)
                markdown.append(f"![暴力破解演示](../{value})")
            else:
                raise ValueError(f"Unknown report element: {kind}")
    target = ROOT / "docs/S-DES实验报告.docx"
    document.save(target)
    (ROOT / "docs/实验报告.md").write_text("\n\n".join(markdown) + "\n", encoding="utf-8")
    return target


def export_pdf(document_path, office_path):
    executable = office_path or shutil.which("soffice") or shutil.which("soffice.exe")
    if not executable:
        raise SystemExit("DOCX and Markdown saved. PDF export needs LibreOffice; supply --office with its executable path.")
    output = ROOT / "output/pdf"
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="sdes_report_") as profile:
        command = [str(executable), "-env:UserInstallation=" + Path(profile).as_uri(),
                   "--headless", "--convert-to", "pdf", "--outdir", str(output), str(document_path)]
        subprocess.run(command, check=True, timeout=120, capture_output=True)
    target = output / (document_path.stem + ".pdf")
    if not target.is_file():
        raise RuntimeError("PDF export did not produce a file")
    return target


def main():
    parser = argparse.ArgumentParser(description="Generate the S-DES experiment report")
    parser.add_argument("--pdf", action="store_true", help="Also export a PDF using LibreOffice")
    parser.add_argument("--office", type=Path, help="Path to the LibreOffice executable")
    args = parser.parse_args()
    evidence = {name: json.loads((ROOT / "evidence" / (name + ".json")).read_text(encoding="utf-8"))
                for name in ("summary", "brute_force", "gui_checks", "cross_vectors", "round_trace")}
    target = write_reports(report_pages(evidence))
    print("Generated Word and Markdown reports")
    if args.pdf:
        export_pdf(target, args.office)
        print("Exported PDF report")


if __name__ == "__main__":
    main()
